"""Convert the Boss workbook into the JSON/JS source used by AIBP.

Run with the bundled Python (openpyxl is only used to read the workbook):
    python tools/export_boss_levels.py path/to/Boss等级属性汇总.xlsx
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HEADERS = (
    "循环", "Boss", "等级", "生命值", "移动（面板）", "移动（生效）", "命中要求",
    "开战晋升", "AT力场加成", "危险／攻击", "危险／命中", "命运／攻击", "命运／命中",
    "闪避骰加成", "新增／移除特质", "条件效果与备注", "属性来源",
)
LEVELS = {label: level for level, label in enumerate(("0", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"))}
BONUS_FIELDS = {
    "开战晋升": "promotions", "AT力场加成": "at", "危险／攻击": "dangerPerAttack",
    "危险／命中": "dangerPerHit", "命运／攻击": "fatePerAttack", "命运／命中": "fatePerHit",
    "闪避骰加成": "evasionDice",
}
# Encrypted source filenames do not expose these two Boss identifiers.
HIDDEN_IDS = {"赫利俄斯（旧日幽魂）": "HELIOS", "黑喙": "BLACKBEAK"}


def integer(value, location: str, *, signed: bool = False) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        raise ValueError(f"{location}: expected an integer, got {value!r}")
    if not signed and value < 0:
        raise ValueError(f"{location}: cannot be negative")
    return int(value)


def read_workbook(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["等级属性"]
        coverage = workbook["说明与覆盖"]
        bosses = {}
        in_coverage = False
        definitions = {}
        for row in coverage.iter_rows(values_only=True):
            if row[0] == "Boss" and row[1] == "已确认等级":
                in_coverage = True
                continue
            if not in_coverage:
                if row[0] and len(row) > 1 and row[1]:
                    definitions[row[0]] = row[1]
                continue
            name, _, cycle, front, back, note = row[:6]
            if not name or not re.fullmatch(r"C\d+", str(cycle)):
                break
            match = re.fullmatch(r"aibp/ps/([A-Z_]+)/[A-Za-z0-9_.-]+\.jpg", front or "")
            boss_id = HIDDEN_IDS.get(name) or (match[1] if match else None)
            if not boss_id or boss_id in bosses or any(boss["name"] == name for boss in bosses.values()):
                raise ValueError(f"Invalid/duplicate Boss identity: {name!r} / {front!r}")
            bosses[boss_id] = {
                "id": boss_id, "name": name, "cycle": str(cycle).lower(),
                "sources": {"front": front, "back": back if back and back.startswith("aibp/") else None},
                "notes": note or "", "levels": {},
            }
        by_name = {boss["name"]: boss for boss in bosses.values()}
        headers = None
        row_count = 0
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if tuple(values[:len(HEADERS)]) == HEADERS:
                headers = list(HEADERS)
                continue
            if not headers or not any(value is not None for value in values):
                continue
            record = dict(zip(headers, values))
            if record["Boss"] not in by_name:
                raise ValueError(f"等级属性 row {row_number}: unknown Boss {record['Boss']!r}")
            boss = by_name[record["Boss"]]
            label = str(record["等级"])
            if label not in LEVELS:
                raise ValueError(f"等级属性 row {row_number}: invalid level {label!r}")
            level = LEVELS[label]
            if str(level) in boss["levels"] or str(record["循环"]).lower() != boss["cycle"]:
                raise ValueError(f"等级属性 row {row_number}: duplicate level or cycle mismatch")
            bonuses = {field: integer(record[header], f"row {row_number} {header}", signed=field == "at")
                       for header, field in BONUS_FIELDS.items()}
            for header in ("移动（面板）", "移动（生效）"):
                value = record[header]
                if value not in ("∞", "—") and not (isinstance(value, str) and re.fullmatch(r"\d+\*", value)):
                    integer(value, f"row {row_number} {header}")
            boss["levels"][str(level)] = {
                "level": level, "label": label,
                "stats": {"wounds": integer(record["生命值"], f"row {row_number} 生命值"),
                          "movementPrinted": record["移动（面板）"], "movement": record["移动（生效）"],
                          "hitRequirement": integer(record["命中要求"], f"row {row_number} 命中要求")},
                "bonuses": bonuses, "traitChanges": record["新增／移除特质"] or "",
                "notes": record["条件效果与备注"] or "",
                "source": {"card": record["属性来源"], "sheet": sheet.title, "row": row_number},
            }
            row_count += 1
        if not headers or not row_count or any(not boss["levels"] for boss in bosses.values()):
            raise ValueError("Workbook is missing headers, rows, or Boss levels")
        return {
            "format": "ato-boss-levels", "schemaVersion": 1,
            "source": {"workbook": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                       "sheet": sheet.title, "bossCount": len(bosses), "levelCount": row_count},
            "bonusSemantics": "Each level stores its printed totals; do not sum earlier levels.",
            "definitions": definitions, "bosses": bosses,
        }
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "aibp")
    args = parser.parse_args()
    data = read_workbook(args.workbook)
    payload = json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "boss-levels-data.json").write_text(payload + "\n", encoding="utf-8")
    javascript = (
        "// Generated from the Boss workbook by tools/export_boss_levels.py.\n"
        "// Edit the workbook, then regenerate both JSON and JS; each level contains totals.\n"
        "(function (root) {\n  \"use strict\";\n  const data = " + payload + ";\n"
        "  root.AIBP_BOSS_LEVEL_DATA = data;\n"
        "  if (typeof module !== \"undefined\" && module.exports) module.exports = data;\n"
        "})(typeof window !== \"undefined\" ? window : globalThis);\n"
    )
    (args.output_dir / "boss-levels-data.js").write_text(javascript, encoding="utf-8")
    print(json.dumps({"bosses": data["source"]["bossCount"], "levels": data["source"]["levelCount"],
                      "output": str(args.output_dir)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
